from pathlib import Path

from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


app = Flask(__name__)
registros = []

CARPETA_EXPORTACIONES = Path(__file__).parent / "exports"
CARPETA_EXPORTACIONES.mkdir(exist_ok=True)

CAMPOS = ["dni", "ubicacion", "region", "provincia", "distrito", "direccion"]


def validar_registro(datos):
    if any(not datos[campo] for campo in CAMPOS):
        return "Completa todos los campos."

    if not datos["dni"].isdigit() or len(datos["dni"]) != 8:
        return "El DNI debe contener exactamente 8 dígitos."

    return None


@app.route("/", methods=["GET", "POST"])
def inicio():
    mensaje_error = None
    mensaje_exito = None
    datos = {campo: "" for campo in CAMPOS}

    if request.method == "POST":
        datos = {
            campo: request.form.get(campo, "").strip()
            for campo in CAMPOS
        }
        mensaje_error = validar_registro(datos)

        if not mensaje_error:
            registros.append(datos.copy())
            mensaje_exito = "Resultado electoral registrado correctamente."
            datos = {campo: "" for campo in CAMPOS}

    return render_template(
        "index.html",
        datos=datos,
        registros=registros,
        mensaje_error=mensaje_error,
        mensaje_exito=mensaje_exito,
    )


@app.route("/exportar")
def exportar_excel():
    if not registros:
        return render_template(
            "index.html",
            datos={campo: "" for campo in CAMPOS},
            registros=registros,
            mensaje_error="Agrega al menos un resultado antes de exportar.",
            mensaje_exito=None,
        ), 400

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados electorales"

    encabezados = [
        "DNI",
        "Ubicación",
        "Región",
        "Provincia",
        "Distrito",
        "Dirección del local de votación",
    ]
    hoja.append(encabezados)

    for registro in registros:
        hoja.append([registro[campo] for campo in CAMPOS])

    relleno = PatternFill("solid", fgColor="0F766E")
    for celda in hoja[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")

    anchos = [14, 24, 20, 20, 20, 42]
    for columna, ancho in zip("ABCDEF", anchos):
        hoja.column_dimensions[columna].width = ancho

    ruta_archivo = CARPETA_EXPORTACIONES / "resultados_electorales.xlsx"
    libro.save(ruta_archivo)

    return send_file(
        ruta_archivo,
        as_attachment=True,
        download_name="resultados_electorales.xlsx",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
